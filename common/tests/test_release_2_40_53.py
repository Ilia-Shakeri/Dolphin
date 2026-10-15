"""2.40.53 — personalization, part one: the company's own fields on customers.

Owner decision 2026-10-14 (Q18): «همگی نیاز است چه بسا بیشتر. بیشتر فیلدهای
پنل باید قابلیت شخصی‌سازی داشته باشند. شخصی‌سازی هم فیچر است و باید فعال
شود.» A field is defined once (`customization.manage`) and appears in the
customer's form, profile, list column, filter, search and spreadsheet; the
values stay on the record whatever happens to the feature or the field.
"""

import io
import json

from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from accounts.models import User
from common.deployment.profile import DeploymentProfile, override_active_profile
from common.deployment.registry import ALL_FEATURES, DEFAULT_OFF_FEATURES
from common.exceptions import BusinessPermissionDenied, BusinessRuleError
from customization import services
from customization.models import CustomFieldDefinition
from sales.models import Customer
from sales.services import create_customer_with_phone, update_customer

PASSWORD = "Strong-pass-953!"


def without_customization():
    return DeploymentProfile(profile_id="client-1", features=frozenset(ALL_FEATURES) - {"customization"}, source="signed-manifest")


class Base(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username="z.admin", password=PASSWORD, role=User.Role.PLATFORM_ADMIN)
        self.manager = User.objects.create_user(username="z.manager", password=PASSWORD, role=User.Role.SALES_MANAGER)
        self.size = services.create_definition(
            actor=self.admin, entity="customer", label="اندازهٔ شرکت", field_type="choice",
            choices=["کوچک", "متوسط", "بزرگ"], show_in_list=True, filterable=True,
        )
        self.since = services.create_definition(actor=self.admin, entity="customer", label="مشتری از", field_type="date")
        self.code = services.create_definition(
            actor=self.admin, entity="customer", label="کد قدیمی", field_type="text", searchable=True, filterable=True,
        )
        self.client_api = APIClient()
        self.client_api.force_authenticate(self.manager)


class DefinitionTests(Base):
    def test_opt_in_and_managed_by_the_capability(self):
        self.assertIn("customization", DEFAULT_OFF_FEATURES)
        self.assertEqual([self.size.key, self.since.key, self.code.key], ["cf_1", "cf_2", "cf_3"])
        with self.assertRaises(BusinessPermissionDenied):
            services.create_definition(actor=self.manager, entity="customer", label="x", field_type="text")
        with self.assertRaises(BusinessRuleError):
            services.create_definition(actor=self.admin, entity="customer", label="اندازهٔ شرکت", field_type="text")
        with self.assertRaises(BusinessRuleError):
            services.create_definition(actor=self.admin, entity="customer", label="رنگ", field_type="choice", choices=[])
        services.move_definition(actor=self.admin, definition=self.code, direction="up")
        order = list(CustomFieldDefinition.objects.order_by("position").values_list("key", flat=True))
        self.assertEqual(order, ["cf_1", "cf_3", "cf_2"])

    def test_the_settings_api_and_tab(self):
        client = APIClient()
        client.force_authenticate(self.admin)
        created = client.post("/api/v1/custom-fields/", {"entity": "customer", "label": "تعداد کارمند", "field_type": "number"}, format="json")
        self.assertEqual(created.status_code, 201, created.content)
        field = next(item for item in created.json()["fields"] if item["label"] == "تعداد کارمند")
        changed = client.patch(f"/api/v1/custom-fields/{field['id']}/", {"is_active": False, "field_type": "text"}, format="json")
        self.assertEqual(changed.status_code, 200)
        after = next(item for item in changed.json()["fields"] if item["id"] == field["id"])
        self.assertEqual((after["is_active"], after["field_type"]), (False, "number"))
        self.assertEqual(self.client_api.get("/api/v1/custom-fields/").status_code, 403)
        with override_active_profile(without_customization()):
            self.assertEqual(client.get("/api/v1/custom-fields/").status_code, 404)
        self.client.force_login(self.admin)
        self.assertContains(self.client.get("/settings/?tab=customization"), 'id="custom-field-form"')


class ValueTests(Base):
    def test_values_are_checked_by_type_and_kept_as_text(self):
        cleaned = services.clean_values("customer", {"cf_1": "بزرگ", "cf_2": "2026-10-15", "cf_3": "  A-12  "})
        self.assertEqual(cleaned, {"cf_1": "بزرگ", "cf_2": "2026-10-15", "cf_3": "A-12"})
        for bad in ({"cf_1": "خیلی بزرگ"}, {"cf_2": "15/10/2026"}, {"cf_9": "x"}):
            with self.assertRaises(BusinessRuleError):
                services.clean_values("customer", bad)
        number = services.create_definition(actor=self.admin, entity="customer", label="عدد", field_type="number")
        self.assertEqual(services.clean_values("customer", {number.key: "۱۲٫۵"}), {number.key: "12.5"})

    def test_the_api_stores_and_merges_and_holds_a_form_to_required_fields(self):
        services.update_definition(actor=self.admin, definition=self.code, required=True)
        refused = self.client_api.post("/api/v1/customers/", {"full_name": "شرکت نمونه", "custom_fields": {"cf_1": "کوچک"}}, format="json")
        self.assertEqual(refused.status_code, 400)
        made = self.client_api.post(
            "/api/v1/customers/", {"full_name": "شرکت نمونه", "custom_fields": {"cf_1": "کوچک", "cf_3": "OLD-7"}}, format="json"
        )
        self.assertEqual(made.status_code, 201, made.content)
        customer = Customer.objects.get(pk=made.json()["id"])
        self.assertEqual(customer.custom_fields, {"cf_1": "کوچک", "cf_3": "OLD-7"})
        patched = self.client_api.patch(f"/api/v1/customers/{customer.pk}/", {"custom_fields": {"cf_2": "2026-01-02"}}, format="json")
        self.assertEqual(patched.status_code, 200, patched.content)
        self.assertEqual(patched.json()["custom_fields"], {"cf_1": "کوچک", "cf_3": "OLD-7", "cf_2": "2026-01-02"})
        # An integration creating a customer is not held to the form's required fields.
        create_customer_with_phone(actor=self.manager, full_name="از سایت")

    def test_without_the_feature_nothing_shows_and_nothing_is_lost(self):
        customer = create_customer_with_phone(actor=self.manager, full_name="الف", custom_fields={"cf_1": "متوسط"})
        with override_active_profile(without_customization()):
            data = self.client_api.get(f"/api/v1/customers/{customer.pk}/").json()
            self.assertNotIn("custom_fields", data)
            with self.assertRaises(BusinessRuleError):
                update_customer(actor=self.manager, customer=customer, custom_fields={"cf_1": "بزرگ"})
            update_customer(actor=self.manager, customer=customer, city="تهران")
        customer.refresh_from_db()
        self.assertEqual(customer.custom_fields, {"cf_1": "متوسط"})
        # A turned-off field keeps its value too.
        services.update_definition(actor=self.admin, definition=self.size, is_active=False)
        update_customer(actor=self.manager, customer=customer, custom_fields={"cf_3": "X"})
        customer.refresh_from_db()
        self.assertEqual(customer.custom_fields, {"cf_1": "متوسط", "cf_3": "X"})


class ListTests(Base):
    def setUp(self):
        super().setUp()
        self.big = create_customer_with_phone(actor=self.manager, full_name="بزرگ‌مهر", custom_fields={"cf_1": "بزرگ", "cf_3": "ZX-101"})
        self.small = create_customer_with_phone(actor=self.manager, full_name="کوچک‌زاده", custom_fields={"cf_1": "کوچک"})

    def names(self, **params):
        response = self.client_api.get("/api/v1/customers/", params)
        self.assertEqual(response.status_code, 200, response.content)
        return {row["full_name"] for row in response.json()["results"]}

    def test_filter_and_search(self):
        self.assertEqual(self.names(custom_fields=json.dumps({"cf_1": "بزرگ"})), {"بزرگ‌مهر"})
        self.assertEqual(self.names(custom_fields=json.dumps({"cf_3": "zx-1"})), {"بزرگ‌مهر"})
        self.assertEqual(self.names(search="ZX-101"), {"بزرگ‌مهر"})
        bad = self.client_api.get("/api/v1/customers/", {"custom_fields": json.dumps({"cf_2": "2026-01-01"})})
        self.assertEqual(bad.status_code, 400)  # not a filterable field

    def test_the_page_shows_columns_filters_and_the_form(self):
        self.client.force_login(self.manager)
        page = self.client.get("/customers/")
        self.assertContains(page, "<th>اندازهٔ شرکت</th>", html=False)
        self.assertContains(page, 'data-custom-filter="cf_1"')
        self.assertContains(page, 'id="custom-field-definitions"')
        profile = self.client.get(f"/customers/{self.big.pk}/")
        self.assertContains(profile, "اندازهٔ شرکت")
        self.assertContains(profile, 'id="edit-customer-custom-fields"')
        with override_active_profile(without_customization()):
            self.assertNotContains(self.client.get("/customers/"), 'id="custom-field-definitions"')

    def test_the_spreadsheet_round_trip(self):
        self.client.force_login(self.manager)
        export = self.client_api.get("/api/v1/exports/customers.xlsx", {"kind": "individual"})
        self.assertEqual(export.status_code, 200)
        sheet = load_workbook(io.BytesIO(b"".join(export.streaming_content) if hasattr(export, "streaming_content") else export.content)).active
        header = [cell.value for cell in sheet[1]]
        self.assertIn("اندازهٔ شرکت", header)
        self.assertIn("بزرگ", [cell.value for cell in sheet[2]] + [cell.value for cell in sheet[3]])
        workbook = Workbook()
        new = workbook.active
        new.append(["full_name", "primary_phone", "اندازهٔ شرکت", "کد قدیمی"])
        new.append(["تازه‌وارد", "09129990071", "متوسط", "NEW-1"])
        new.append(["بد", "09129990072", "غول", ""])
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        stream.name = "customers.xlsx"
        admin = APIClient()
        admin.force_authenticate(self.admin)
        result = admin.post("/api/v1/customers/import-xlsx/", {"file": stream, "kind": "individual"}, format="multipart")
        self.assertEqual(result.status_code, 200, result.content)
        self.assertEqual((result.json()["created"], result.json()["invalid"]), (1, 1))
        self.assertEqual(Customer.objects.get(full_name="تازه‌وارد").custom_fields, {"cf_1": "متوسط", "cf_3": "NEW-1"})
