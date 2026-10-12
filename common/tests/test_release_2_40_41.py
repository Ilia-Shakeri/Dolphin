"""Roadmap «اولویت ۱.۵» (2.40.41).

The ranking as a workbook, a monthly sales target for each marketer, moving
several parcels to one postal status at once, and detail pages that redraw in
place when their record changes elsewhere.
"""

import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.access import has_any_capability
from accounts.models import User
from auditlog.models import ActivityLog
from billing.services import create_invoice, issue_invoice
from common import jalali
from reports.targets import build_monthly_targets, current_month, month_window
from reports.xlsx import MARKETER_HEADERS
from sales.models import SalesTarget
from sales.services import (
    create_customer_with_phone,
    create_product,
    deactivate_sales_document,
    register_sales_document,
)

ROOT = Path(__file__).resolve().parents[2]
JS = ROOT / "common" / "static" / "common" / "js"
PASSWORD = "Strong-pass-741!"


class _Marketers(TestCase):
    def setUp(self):
        self.manager = User.objects.create_user(username="t.manager", password=PASSWORD, role=User.Role.SALES_MANAGER)
        self.it = User.objects.create_user(username="t.it", password=PASSWORD, role=User.Role.COMPANY_IT)
        self.first = User.objects.create_user(username="t.first", password=PASSWORD, role=User.Role.SALES_AGENT)
        self.second = User.objects.create_user(username="t.second", password=PASSWORD, role=User.Role.SALES_AGENT)
        self.product = create_product(actor=self.manager, sku="T-1", name="کالا", current_price=Decimal("100.00"))

    def sell(self, agent, phone, quantity):
        customer = create_customer_with_phone(actor=agent, full_name=f"مشتری {phone}", phone={"raw_phone": phone})
        invoice = create_invoice(actor=agent, customer=customer, items=[
            {"product": self.product, "quantity": quantity, "unit_price": self.product.current_price},
        ])
        return issue_invoice(actor=self.manager, invoice=invoice)

    def client_for(self, user):
        client = APIClient()
        client.force_login(user)
        return client


class RankingExportTests(_Marketers):
    def window(self):
        now = timezone.now()
        return {"period_start": (now - datetime.timedelta(days=1)).isoformat(), "period_end": (now + datetime.timedelta(minutes=1)).isoformat()}

    def test_the_ranking_downloads_as_the_page_shows_it(self):
        self.sell(self.first, "09121110301", 5)
        response = self.client_for(self.manager).get("/api/v1/exports/marketers.xlsx", self.window())
        self.assertEqual(response.status_code, 200, response.content[:300])
        self.assertIn("spreadsheetml", response["Content-Type"])
        book = load_workbook(BytesIO(response.content))
        rows = list(book["marketers"].iter_rows(values_only=True))
        self.assertEqual(rows[0], MARKETER_HEADERS)
        self.assertEqual((rows[1][0], rows[1][3]), (1, "t.first"))
        summary = dict(book["summary"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["ordering"], "sales_amount")
        self.assertEqual(summary["marketers"], 2)

    def test_the_export_keeps_the_rankings_gate(self):
        response = self.client_for(self.first).get("/api/v1/exports/marketers.xlsx", self.window())
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response["Content-Type"], "application/json")


class MonthlyTargetTests(_Marketers):
    def month(self):
        return current_month(timezone.now())

    def test_a_manager_sets_targets_and_sees_progress(self):
        invoice = self.sell(self.first, "09121110311", 5)
        year, month = self.month()
        client = self.client_for(self.manager)
        response = client.put("/api/v1/reports/marketers/targets/", {
            "year": year, "month": month,
            "targets": [{"user_id": self.first.pk, "amount": "1000.00"}, {"user_id": self.second.pk, "amount": None}],
        }, format="json")
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()
        self.assertTrue(data["can_manage"])
        rows = {row["user_id"]: row for row in data["results"]}
        self.assertEqual(Decimal(rows[self.first.pk]["target"]), Decimal("1000.00"))
        self.assertEqual(Decimal(rows[self.first.pk]["achieved"]), invoice.total_amount)
        self.assertEqual(Decimal(rows[self.first.pk]["percent"]), (invoice.total_amount / 10).quantize(Decimal("0.1")))
        self.assertIsNone(rows[self.second.pk]["target"])
        self.assertEqual(data["results"][0]["user_id"], self.first.pk)  # those with a target first
        self.assertTrue(ActivityLog.objects.filter(operation="sales_target.set", object_type="sales.salestarget").exists())

    def test_an_empty_amount_clears_a_target(self):
        year, month = self.month()
        SalesTarget.objects.create(user=self.first, year=year, month=month, amount=Decimal("50.00"), set_by=self.manager)
        response = self.client_for(self.manager).put("/api/v1/reports/marketers/targets/", {
            "year": year, "month": month, "targets": [{"user_id": self.first.pk, "amount": ""}],
        }, format="json")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertFalse(SalesTarget.objects.filter(user=self.first).exists())
        self.assertTrue(ActivityLog.objects.filter(operation="sales_target.cleared").exists())

    def test_reading_needs_company_reports_and_writing_needs_the_target_capability(self):
        year, month = self.month()
        body = {"year": year, "month": month, "targets": [{"user_id": self.first.pk, "amount": "10"}]}
        self.assertEqual(self.client_for(self.first).get("/api/v1/reports/marketers/targets/").status_code, 403)
        it = self.client_for(self.it)
        self.assertEqual(it.get("/api/v1/reports/marketers/targets/").status_code, 200)
        self.assertFalse(it.get("/api/v1/reports/marketers/targets/").json()["can_manage"])
        self.assertEqual(it.put("/api/v1/reports/marketers/targets/", body, format="json").status_code, 403)
        self.assertFalse(SalesTarget.objects.exists())
        self.assertTrue(has_any_capability(self.manager, "sales_targets.manage"))

    def test_only_active_marketers_take_a_target(self):
        year, month = self.month()
        response = self.client_for(self.manager).put("/api/v1/reports/marketers/targets/", {
            "year": year, "month": month, "targets": [{"user_id": self.manager.pk, "amount": "10"}],
        }, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(SalesTarget.objects.exists())

    def test_a_month_that_does_not_exist_is_refused(self):
        response = self.client_for(self.manager).get("/api/v1/reports/marketers/targets/", {"year": 1405, "month": 13})
        self.assertEqual(response.status_code, 400)

    def test_the_month_is_counted_in_tehran_time(self):
        start, end = month_window(1405, 7)
        self.assertEqual(start.date(), jalali.from_jalali(1405, 7, 1))
        self.assertEqual(end.date(), jalali.from_jalali(1405, 8, 1))
        self.assertEqual(start.utcoffset(), datetime.timedelta(hours=3, minutes=30))
        # The last month of the year wraps to Farvardin of the next.
        self.assertEqual(month_window(1405, 12)[1].date(), jalali.from_jalali(1406, 1, 1))
        report = build_monthly_targets(actor=self.manager, year=1405, month=7)
        self.assertEqual(report["label"], "مهر ۱۴۰۵")


class BulkPostalStatusTests(TestCase):
    def setUp(self):
        self.manager = User.objects.create_user(username="b.manager", password=PASSWORD, role=User.Role.SALES_MANAGER)
        customer = create_customer_with_phone(actor=self.manager, full_name="مشتری پستی", province="تهران", city="تهران")
        self.documents = [
            register_sales_document(actor=self.manager, customer=customer, sale=None, document_number=f"B-{index}", postal_status="in_store")
            for index in range(3)
        ]
        self.client = APIClient()
        self.client.force_login(self.manager)

    def test_the_chosen_parcels_move_together_and_the_rest_are_named(self):
        moved, same, inactive = self.documents
        same_status = "handed_to_post"
        self.client.post(f"/api/v1/sales-documents/{same.pk}/transition-postal-status/", {"to_status": same_status}, format="json")
        deactivate_sales_document(actor=self.manager, document=inactive)
        response = self.client.post("/api/v1/sales-documents/bulk-transition-postal-status/", {
            "ids": [moved.pk, same.pk, inactive.pk, 999999], "to_status": same_status, "reason": "تحویل گروهی",
        }, format="json")
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()
        self.assertEqual(data["changed"], [moved.pk])
        self.assertEqual({(row["id"], row["reason"]) for row in data["skipped"]},
                         {(same.pk, "same_status"), (inactive.pk, "inactive"), (999999, "not_found")})
        moved.refresh_from_db()
        self.assertEqual(moved.postal_status, same_status)
        self.assertEqual(moved.postal_history.first().reason, "تحویل گروهی")
        self.assertTrue(ActivityLog.objects.filter(operation="sales_document.postal_status_changed", object_id=str(moved.pk)).exists())

    def test_a_reader_without_manage_cannot_move_any(self):
        reader = User.objects.create_user(username="b.agent", password=PASSWORD, role=User.Role.SALES_AGENT)
        client = APIClient()
        client.force_login(reader)
        response = client.post("/api/v1/sales-documents/bulk-transition-postal-status/", {
            "ids": [self.documents[0].pk], "to_status": "handed_to_post",
        }, format="json")
        self.assertIn(response.status_code, {403, 404})
        self.documents[0].refresh_from_db()
        self.assertEqual(self.documents[0].postal_status, "in_store")

    def test_at_most_two_hundred_at_once(self):
        response = self.client.post("/api/v1/sales-documents/bulk-transition-postal-status/", {
            "ids": list(range(1, 202)), "to_status": "handed_to_post",
        }, format="json")
        self.assertEqual(response.status_code, 400)


class LiveRecordScriptTests(SimpleTestCase):
    DETAIL_MODULES = (
        "billing/invoice-detail.js", "billing/order-detail.js", "billing/payment-detail.js",
        "leads/lead-detail.js", "sales/sale-detail.js", "sales/sales-document-detail.js",
        "interactions/interaction-detail.js", "after_sales/after-sales-detail.js",
        "campaigns/campaign-detail.js", "products/product-detail.js",
        "products/product-category-detail.js", "inventory/warehouse-detail.js",
        "customers/person-profile.js",
    )

    def test_every_detail_page_registers_its_record(self):
        for name in self.DETAIL_MODULES:
            with self.subTest(module=name):
                self.assertIn("liveRecord({", (JS / "features" / name).read_text(encoding="utf-8"))

    def test_a_registered_page_redraws_instead_of_reloading(self):
        page = (JS / "shell" / "live-page.js").read_text(encoding="utf-8")
        self.assertIn("if (liveRecordFor()) {", page)
        self.assertIn('showChangedNotice(record)', page)
        record = (JS / "ui" / "live-record.js").read_text(encoding="utf-8")
        self.assertIn('if (edited) return "edited";', record)
